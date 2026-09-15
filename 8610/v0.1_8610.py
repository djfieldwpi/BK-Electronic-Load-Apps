# pyinstaller --onefile --windowed --name="BatteryTester8610" v0.1_8610.py

from PyQt6.QtWidgets import (QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QLabel,
                             QLineEdit, QPushButton, QStyle, QMessageBox)
from PyQt6.QtCore import QSize, QTimer
from PyQt6.QtGui import QDoubleValidator

from openpyxl import Workbook, load_workbook
from datetime import datetime

from pathlib import Path

import pyvisa, time, json

CONFIG_FILE = Path.home() / "BatteryTester" / "8610_settings.json"
CONFIG_FILE.parent.mkdir(exist_ok=True)

def load_settings():
    global defaultPath, defaultCurrLimit, defaultLimitDelay, defaultConstPower, defaultStopVoltage, defaultMonitorTime, defaultRecordTime

    if not CONFIG_FILE.exists():
        return

    with open(CONFIG_FILE, "r") as f:
        data = json.load(f)

    defaultPath = data.get("defaultPath", "")
    defaultCurrLimit = data.get("defaultCurrLimit", 0.0)
    defaultLimitDelay = data.get("defaultLimitDelay", 0.0)
    defaultConstPower = data.get("defaultConstPower", 0.0)
    defaultStopVoltage = data.get("defaultStopVoltage", 0.0)
    defaultMonitorTime = data.get("defaultMonitorTime", 1.0)
    defaultRecordTime = data.get("defaultRecordTime", 1.0)

def save_settings():
    data = {
        "defaultPath": defaultPath,
        "defaultCurrLimit": defaultCurrLimit,
        "defaultLimitDelay": defaultLimitDelay,
        "defaultConstPower": defaultConstPower,
        "defaultStopVoltage": defaultStopVoltage,
        "defaultMonitorTime": defaultMonitorTime,
        "defaultRecordTime": defaultRecordTime
    }

    with open(CONFIG_FILE, "w") as f:
        json.dump(data, f, indent=4)

defaultPath = ""
defaultCurrLimit = 0.0
defaultLimitDelay = 0.0
defaultConstPower = 0.0
defaultStopVoltage = 0.0
defaultMonitorTime = 1.0
defaultRecordTime = 1.0

class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()

        # Misc
        self.started = False
        self.testTimer = QTimer(self)
        self.testTimer.timeout.connect(self.testLoop)
        self.lastMonitor = 0.0
        self.lastRecord = 0.0

        # Connection
        self.rm = None
        self.inst = None

        # Monitor
        self.monitorVoltage = 0.0
        self.monitorCurrent = 0.0
        self.monitorPower = 0.0

        # Record
        self.recordVoltage = 0.0
        self.recordCurrent = 0.0
        self.recordPower = 0.0
        self.recordStatus = ""
        self.recordTimer = 0.0

        # Excel
        self.wb = None
        self.filePath = ""

        self.setWindowTitle("Battery Load Tester")
        self.setFixedSize(QSize(600, 550))

        # Overall Layout
        vLayout = QVBoxLayout()
        fileLayout = QHBoxLayout()
        limitLayout = QHBoxLayout()
        delayLayout = QHBoxLayout()
        constLayout = QHBoxLayout()
        stopLayout = QHBoxLayout()
        buttonLayout = QHBoxLayout()
        monitorLayout = QHBoxLayout()
        recordLayout = QHBoxLayout()
        statusLayout = QHBoxLayout()
        voltageLayout = QHBoxLayout()
        currentLayout = QHBoxLayout()
        powerLayout = QHBoxLayout()
        timerLayout = QHBoxLayout()
        vLayout.addLayout(fileLayout)
        vLayout.addLayout(limitLayout)
        vLayout.addLayout(delayLayout)
        vLayout.addLayout(constLayout)
        vLayout.addLayout(stopLayout)
        vLayout.addLayout(monitorLayout)
        vLayout.addLayout(recordLayout)
        vLayout.addLayout(buttonLayout)
        vLayout.addLayout(statusLayout)
        vLayout.addLayout(voltageLayout)
        vLayout.addLayout(currentLayout)
        vLayout.addLayout(powerLayout)
        vLayout.addLayout(timerLayout)

        # ========== SETTINGS ==========

        # File Layout
        self.fileLabel = QLabel("<b style=\"font-size: 20px\">File Name:</b>")
        self.fileEntry = QLineEdit()
        self.fileEntry.setStyleSheet("""
                                    QLineEdit{
                                    font-size: 20px
                                    }
                                    """)
        self.fileEntry.setText(defaultPath)
        self.fileEntry.setEnabled(not self.started)
        self.testLabel = QLabel("<b style=\"font-size: 20px\">Test ID:</b>")
        self.testEntry = QLineEdit()
        self.testEntry.setStyleSheet("""
                                    QLineEdit{
                                    font-size: 20px
                                    }
                                    """)
        self.testEntry.setPlaceholderText("Test ID...")
        self.testEntry.setMaximumWidth(200)
        self.testEntry.setEnabled(not self.started)

        fileLayout.addWidget(self.fileLabel)
        fileLayout.addWidget(self.fileEntry)
        fileLayout.addWidget(self.testLabel)
        fileLayout.addWidget(self.testEntry)

        # Current Limit Layout
        self.limitLabel = QLabel("<b style=\"font-size: 20px\">Current Limit:</b>")
        self.limitEntry = QLineEdit()
        self.limitEntry.setStyleSheet("""
                                    QLineEdit{
                                    font-size: 20px
                                    }
                                    """)
        self.limitEntry.setPlaceholderText(str(defaultCurrLimit) + " A")
        self.limitEntry.setValidator(QDoubleValidator())
        self.limitEntry.setEnabled(not self.started)
        self.limitEntry.setMaximumWidth(200)

        limitLayout.addWidget(self.limitLabel)
        limitLayout.addWidget(self.limitEntry)

        # Limit Delay Layout
        self.delayLabel = QLabel("<b style=\"font-size: 20px\">Limit Protection Delay:</b>")
        self.delayEntry = QLineEdit()
        self.delayEntry.setStyleSheet("""
                                    QLineEdit{
                                    font-size: 20px
                                    }
                                    """)
        self.delayEntry.setPlaceholderText(str(defaultLimitDelay) + " s")
        self.delayEntry.setValidator(QDoubleValidator())
        self.delayEntry.setEnabled(not self.started)
        self.delayEntry.setMaximumWidth(200)

        delayLayout.addWidget(self.delayLabel)
        delayLayout.addWidget(self.delayEntry)

        # Constant Power Layout
        self.constLabel = QLabel("<b style=\"font-size: 20px\">Constant Power:</b>")
        self.constEntry = QLineEdit()
        self.constEntry.setStyleSheet("""
                                    QLineEdit{
                                    font-size: 20px
                                    }
                                    """)
        self.constEntry.setPlaceholderText(str(defaultConstPower) + " W")
        self.constEntry.setValidator(QDoubleValidator())
        self.constEntry.setEnabled(not self.started)
        self.constEntry.setMaximumWidth(200)

        constLayout.addWidget(self.constLabel)
        constLayout.addWidget(self.constEntry)

        # Stop Voltage Layout
        self.stopLabel = QLabel("<b style=\"font-size: 20px\">Stop Voltage:</b>")
        self.stopEntry = QLineEdit()
        self.stopEntry.setStyleSheet("""
                                    QLineEdit{
                                    font-size: 20px
                                    }
                                    """)
        self.stopEntry.setPlaceholderText(str(defaultStopVoltage) + " V")
        self.stopEntry.setValidator(QDoubleValidator())
        self.stopEntry.setEnabled(not self.started)
        self.stopEntry.setMaximumWidth(200)

        stopLayout.addWidget(self.stopLabel)
        stopLayout.addWidget(self.stopEntry)

        # Monitor Time Layout
        self.monitorLabel = QLabel("<b style=\"font-size: 20px\">Monitor Time:</b>")
        self.monitorEntry = QLineEdit()
        self.monitorEntry.setStyleSheet("""
                                    QLineEdit{
                                    font-size: 20px
                                    }
                                    """)
        self.monitorEntry.setPlaceholderText(str(defaultMonitorTime) + " s")
        self.monitorEntry.setValidator(QDoubleValidator())
        self.monitorEntry.setEnabled(not self.started)
        self.monitorEntry.setMaximumWidth(200)

        monitorLayout.addWidget(self.monitorLabel)
        monitorLayout.addWidget(self.monitorEntry)

        # Record Time Layout
        self.recordLabel = QLabel("<b style=\"font-size: 20px\">Record Time:</b>")
        self.recordEntry = QLineEdit()
        self.recordEntry.setStyleSheet("""
                                    QLineEdit{
                                    font-size: 20px
                                    }
                                    """)
        self.recordEntry.setPlaceholderText(str(defaultRecordTime) + " s")
        self.recordEntry.setValidator(QDoubleValidator())
        self.recordEntry.setEnabled(not self.started)
        self.recordEntry.setMaximumWidth(200)

        recordLayout.addWidget(self.recordLabel)
        recordLayout.addWidget(self.recordEntry)

        # ========== BUTTONS ==========

        # Button Layout
        self.startButton = QPushButton("START TEST")
        self.startButton.setEnabled(not self.started)
        self.startButton.clicked.connect(self.start)
        self.startButton.setStyleSheet("""
                                            QPushButton {
                                            background-color: #00BB00;
                                            font-size: 35px;
                                            }
                                            QPushButton:hover {
                                            background-color: #00CB00;
                                            }
                                            QPushButton:pressed {
                                            background-color: #009B00;
                                            }
                                            """)
        
        self.endButton = QPushButton("END TEST")
        self.endButton.setEnabled(self.started)
        self.endButton.clicked.connect(self.end)
        self.endButton.setStyleSheet("""
                                        QPushButton {
                                        background-color: #BB0000;
                                       font-size: 35px;
                                        }
                                        QPushButton:hover {
                                       background-color: #CB0000;
                                        }
                                        QPushButton:pressed {
                                        background-color: #9B0000;
                                        }
                                        """)
        
        buttonLayout.addWidget(self.startButton)
        buttonLayout.addWidget(self.endButton)

        # ========== DISPLAY ==========

        # Status Layout 
        self.statusLabel = QLabel(f"<b style=\"font-size: 20px;\">Enter test number to begin testing.</b></div>")

        statusLayout.addWidget(self.statusLabel)
        
        # Voltage Layout
        self.voltageLabel = QLabel(f"<b style=\"font-size: 20px;\">Voltage: {str(self.monitorVoltage)} V</b></div>")

        voltageLayout.addWidget(self.voltageLabel)

        # Current Layout
        self.currentLabel = QLabel(f"<b style=\"font-size: 20px;\">Current: {str(self.monitorCurrent)} A</b></div>")

        currentLayout.addWidget(self.currentLabel)

        # Power Layout
        self.powerLabel = QLabel(f"<b style=\"font-size: 20px;\">Power: {str(self.monitorPower)} W</b></div>")

        powerLayout.addWidget(self.powerLabel)

        # Timer Layout
        self.timerLabel = QLabel(f"<b style=\"font-size: 20px;\">Time until record: {str(self.recordTimer)} s</b></div>")

        timerLayout.addWidget(self.timerLabel)

        window = QWidget()
        window.setLayout(vLayout)
        self.setCentralWidget(window)

    def updateUI(self):
        self.fileEntry.setEnabled(not self.started)
        self.testEntry.setEnabled(not self.started)
        self.limitEntry.setEnabled(not self.started)
        self.delayEntry.setEnabled(not self.started)
        self.constEntry.setEnabled(not self.started)
        self.stopEntry.setEnabled(not self.started)
        self.monitorEntry.setEnabled(not self.started)
        self.recordEntry.setEnabled(not self.started)

        self.startButton.setEnabled(not self.started)
        self.endButton.setEnabled(self.started)

        self.statusLabel.setText(f"<b style=\"font-size: 20px;\">Status: {self.recordStatus}</b></div>")
        self.voltageLabel.setText(f"<b style=\"font-size: 20px;\">Voltage: {str(self.monitorVoltage)} V</b></div>")
        self.currentLabel.setText(f"<b style=\"font-size: 20px;\">Current: {str(self.monitorCurrent)} A</b></div>")
        self.powerLabel.setText(f"<b style=\"font-size: 20px;\">Power: {str(self.monitorPower)} W</b></div>")
        self.timerLabel.setText(f"<b style=\"font-size: 20px;\">Time until record: {self.recordTimer:.3f} s</b></div>")


    def connect(self):
        self.rm = pyvisa.ResourceManager()
        resources = self.rm.list_resources()

        if not resources:
            return False
    
        self.inst = self.rm.open_resource(resources[0])
        return True

    def monitor(self, record):
        status = int(self.inst.query("STAT:QUES:COND?"))
        print(status)

        oc = bool(status & (1 << 1))
        ps = bool(status & (1 << 4))

        if oc and ps:
            self.recordVoltage = self.monitorVoltage
            self.recordCurrent = self.monitorCurrent
            self.recordPower = self.monitorPower
            self.record()
            self.recordStatus = "Current Protection Exceeded"
            return False

        else:
            v = float(self.inst.query("MEAS:VOLT?"))
            c = float(self.inst.query("MEAS:CURR?"))
            p = float(self.inst.query("FETC:POW?"))

            print([v, c, p])

            if v < defaultStopVoltage:
                self.recordVoltage = self.monitorVoltage
                self.recordCurrent = self.monitorCurrent
                self.recordPower = self.monitorPower
                self.record()
                self.recordStatus = "Voltage below Stop Threshold"
                return False

            self.monitorVoltage = v
            self.monitorCurrent = c
            self.monitorPower = p

            if record:
                self.recordVoltage = v
                self.recordCurrent = c
                self.recordPower = p
                self.record()
        
        return True

    def initializeWorkbook(self):
        self.wb = Workbook()
        sheet = self.wb.active
        sheet.title = self.testEntry.text()
        sheet["A1"] = "Test ID"
        sheet["B1"] = "Current Limit"
        sheet["C1"] = "Constant Power"
        sheet["D1"] = "Stop Voltage"
        sheet["E1"] = "Time"
        sheet["F1"] = "Voltage"
        sheet["G1"] = "Current"
        sheet["H1"] = "Power"
        sheet["I1"] = "Status"
        self.wb.save(self.filePath)

    def record(self):
        sheet = self.wb.active
        sheet.append([
            self.testEntry.text(),
            float(defaultCurrLimit),
            float(defaultConstPower),
            float(defaultStopVoltage),
            datetime.now(),
            float(self.recordVoltage),
            float(self.recordCurrent),
            float(self.recordPower),
            str(self.recordStatus)
        ])
        self.wb.save(self.filePath)

    def setDefaults(self):
        global defaultPath, defaultCurrLimit, defaultLimitDelay, defaultConstPower, defaultStopVoltage, defaultMonitorTime, defaultRecordTime

        success = True

        if self.fileEntry.text().strip():
            defaultPath = self.fileEntry.text()

        if self.limitEntry.text().strip():
            defaultCurrLimit = float(self.limitEntry.text())
        if self.delayEntry.text().strip():
            temp = float(self.delayEntry.text())
            if temp > 60.0:
                self.delayEntry.clear()
                self.delayEntry.setPlaceholderText("delay too long")
                success = False
            else:
                defaultLimitDelay = temp
        if self.constEntry.text().strip():
            defaultConstPower = float(self.constEntry.text())
        if self.stopEntry.text().strip():
            defaultStopVoltage = float(self.stopEntry.text())
        if self.monitorEntry.text().strip():
            temp = float(self.monitorEntry.text())
            if temp > 0.0:
                defaultMonitorTime = temp
            else:
                self.monitorEntry.clear()
                self.monitorEntry.setPlaceholderText("time too short")
                success = False
        if self.recordEntry.text().strip():
            temp = float(self.recordEntry.text())
            if temp > defaultMonitorTime:
                defaultRecordTime = temp
            else:
                self.recordEntry.clear()
                self.recordEntry.setPlaceholderText("time too short")
                success = False

        save_settings()
        return success

    def start(self):
        self.started = True
        self.updateUI()

        if not self.setDefaults():
            self.started = False
            self.updateUI()
            return

        if not self.testEntry.text().strip():
            self.testEntry.clear()
            self.testEntry.setPlaceholderText("ID REQUIRED")
            self.started = False
            self.updateUI()
            return

        self.filePath = defaultPath + "\\test" + self.testEntry.text() + ".xlsx"
        folder = Path(defaultPath)

        if not folder.exists():
            folder.mkdir(parents = True, exist_ok=True)

        if Path(self.filePath).exists():
                    self.testEntry.clear()
                    self.testEntry.setPlaceholderText("Already Exists.")
                    self.started = False
                    self.updateUI()
                    return
        else:
            self.initializeWorkbook()
        
        if not self.connect():
            # Error Message
            self.started = False
            self.updateUI()
            return
        
        self.inst.write("CURR:PROT:STAT ON")
        self.inst.write("CURR:PROT:LEV " + str(defaultCurrLimit))
        self.inst.write("CURR:PROT:DEL " + str(defaultLimitDelay))

        self.inst.write("FUNC POW")
        self.inst.write("POW " + str(defaultConstPower))

        self.inst.write("INP ON")

        self.lastMonitor = time.monotonic()
        self.lastRecord = time.monotonic()

        self.testTimer.start(50)

    def testLoop(self):
        now = time.monotonic()

        self.recordTimer = max(0.0, defaultRecordTime - (now - self.lastRecord))

        if now - self.lastMonitor >= defaultMonitorTime:

            shouldRecord = now - self.lastRecord >= defaultRecordTime
            flag = self.monitor(shouldRecord)
            self.lastMonitor = now

            if shouldRecord:
                self.lastRecord = now

            if not flag:
                self.testTimer.stop()
                self.end()
                return
        
        self.recordStatus = "Testing..."
        self.updateUI()

    def end(self):
        self.wb.close()
        self.testTimer.stop()
        self.inst.write("POW 0")
        self.inst.write("INP OFF")
        self.inst.write("SYST:LOC")
        self.inst.close()
        self.rm.close()
        self.started = False
        self.updateUI()

    def closeEvent(self, event):
        if self.started:
            msg = QMessageBox(self)
            msg.setWindowTitle("End Test?")
            msg.setText("Close and save test?")

            saveButton = msg.addButton("Save Test", QMessageBox.ButtonRole.AcceptRole)
            cancelButton = msg.addButton("Cancel", QMessageBox.ButtonRole.RejectRole)

            msg.exec()

            if msg.clickedButton() == saveButton:
                self.end()
                save_settings()
                event.accept()
            elif msg.clickedButton() == cancelButton:
                event.ignore()
        else:
            self.setDefaults()
            save_settings()
            event.accept()

load_settings()
app = QApplication([])
window = MainWindow()
window.show()
app.exec()